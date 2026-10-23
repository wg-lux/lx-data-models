"""Lossless, bounded Excel transport for the canonical study exchange bundle."""

import json
from hashlib import sha256
from io import BytesIO
from itertools import zip_longest
from typing import cast
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

from lx_dtypes.models.contracts.study_exchange import StudyExchangeBundle
from lx_dtypes.package_archive import read_bounded_zip

MAX_EXCHANGE_BYTES = 32 * 1024 * 1024
_CHUNK_SIZE = 16000  # Also below Excel's 32767 UTF-16-unit limit for astral text.


def parse_study_exchange(
    value: StudyExchangeBundle | str | bytes,
) -> StudyExchangeBundle:
    """Revalidate mutable nested objects even when supplied as an existing model."""
    encoded = (
        value.model_dump_json() if isinstance(value, StudyExchangeBundle) else value
    )
    if (
        len(encoded.encode() if isinstance(encoded, str) else encoded)
        > MAX_EXCHANGE_BYTES
    ):
        raise ValueError("study exchange exceeds the size limit")
    return StudyExchangeBundle.model_validate_json(encoded)


def study_exchange_digest(bundle: StudyExchangeBundle) -> str:
    """Stable transport hash of the complete canonical evidence."""
    bundle = parse_study_exchange(bundle)
    payload = json.dumps(
        bundle.model_dump(mode="json"),
        sort_keys=True,
        separators=(",", ":"),
        ensure_ascii=False,
    )
    return sha256(payload.encode()).hexdigest()


def _row(sheet: Worksheet, values: list[str | int | float | bool | None]) -> None:
    if sheet.max_row >= 1_048_576:
        raise ValueError("study exchange exceeds Excel's row limit")
    row_number = (
        sheet.max_row + 1
        if sheet.max_row > 1 or sheet.cell(1, 1).value is not None
        else 1
    )
    for column, value in enumerate(values, 1):
        if isinstance(value, str) and len(value.encode("utf-16-le")) // 2 > 32767:
            raise ValueError("study exchange exceeds Excel's cell limit")
        cell = cast(Cell, sheet.cell(row=row_number, column=column))
        cell.value = value
        if isinstance(value, str):
            cell.data_type = "s"  # Literal text, including strings beginning with '='.


def study_exchange_to_xlsx(bundle: StudyExchangeBundle) -> bytes:
    """Readable projections plus canonical chunks; projections are not import inputs."""
    bundle = parse_study_exchange(bundle)
    workbook = Workbook()
    sheet = workbook.active
    assert isinstance(sheet, Worksheet)
    sheet.title = "Collections"
    _row(
        sheet,
        [
            "record_uuid",
            "study_uuid",
            "cohort_uuid",
            "examination_uuid",
            "patient_uuid",
            "examination",
            "module",
            "version",
            "evaluation_ok",
            "evaluated_at",
        ],
    )
    observations = workbook.create_sheet("Observations")
    _row(
        observations,
        [
            "record_uuid",
            "finding_uuid",
            "finding",
            "classification_group_uuid",
            "choice_uuid",
            "classification",
            "choice",
            "descriptor_uuid",
            "descriptor",
            "value_type",
            "unit",
            "value",
            "value_json",
        ],
    )
    descriptors_by_snapshot = {
        package.snapshot_id: {
            descriptor.name: descriptor
            for descriptor in package.records["ClassificationChoiceDescriptor"]
        }
        for package in bundle.packages
    }
    for item in bundle.collections:
        _row(
            sheet,
            [
                str(item.record_uuid),
                str(item.study_uuid),
                str(item.cohort_uuid or ""),
                str(item.patient_examination_uuid),
                item.record.patient,
                item.record.examination,
                item.record.knowledge_base_module,
                item.record.knowledge_base_version,
                item.evaluation.ok,
                item.evaluated_at.isoformat(),
            ],
        )
        for finding in item.record.patient_findings:
            base: list[str | int | float | bool | None] = [
                str(item.record_uuid),
                str(finding.uuid),
                finding.finding,
            ]
            if not finding.patient_finding_classifications:
                _row(observations, base + [None] * 10)
            for group in finding.patient_finding_classifications:
                if not group.patient_finding_classification_choices:
                    _row(observations, base + [str(group.uuid)] + [None] * 9)
                for choice in group.patient_finding_classification_choices:
                    prefix = base + [
                        str(group.uuid),
                        str(choice.uuid),
                        choice.classification,
                        choice.classification_choice,
                    ]
                    descriptors = (
                        choice.patient_finding_classification_choice_descriptors
                    )
                    if not descriptors:
                        _row(observations, prefix + [None] * 6)
                    for descriptor in descriptors:
                        definition = descriptors_by_snapshot[item.package_snapshot_id][
                            descriptor.classification_choice_descriptor
                        ]
                        value = descriptor.descriptor_value
                        _row(
                            observations,
                            prefix
                            + [
                                str(descriptor.uuid),
                                descriptor.classification_choice_descriptor,
                                type(value).__name__,
                                str(getattr(definition, "unit", "")),
                                None if isinstance(value, list) else value,
                                json.dumps(value, ensure_ascii=False),
                            ],
                        )
    definitions = workbook.create_sheet("Definitions")
    _row(
        definitions,
        ["kind", "uuid", "study_uuid", "name", "hypothesis", "definition_json"],
    )
    for study in bundle.studies:
        _row(
            definitions,
            [
                "study",
                str(study.study_uuid),
                str(study.study_uuid),
                study.definition.name,
                study.definition.hypothesis,
                study.definition.model_dump_json(),
            ],
        )
    for cohort in bundle.cohorts:
        _row(
            definitions,
            [
                "cohort",
                str(cohort.cohort_uuid),
                str(cohort.study_uuid),
                cohort.definition.name,
                cohort.definition.hypothesis,
                cohort.definition.model_dump_json(),
            ],
        )
    evaluations = workbook.create_sheet("Evaluation Issues")
    _row(evaluations, ["record_uuid", "source", "level", "code", "message"])
    for item in bundle.collections:
        for issue in item.evaluation.issues:
            _row(
                evaluations,
                [
                    str(item.record_uuid),
                    issue.source,
                    issue.level,
                    issue.code,
                    issue.message,
                ],
            )
    canonical = workbook.create_sheet("LX Canonical")
    _row(canonical, ["study_exchange_v1", "canonical_json_chunks"])
    encoded = bundle.model_dump_json()
    for offset in range(0, len(encoded), _CHUNK_SIZE):
        _row(canonical, [offset // _CHUNK_SIZE, encoded[offset : offset + _CHUNK_SIZE]])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def study_exchange_from_xlsx(data: bytes) -> StudyExchangeBundle:
    """Restore canonical evidence and reject edited display projections."""
    # Bound all decompressed XML before openpyxl parses it; never extract paths.
    members = read_bounded_zip(
        BytesIO(data),
        max_zip_bytes=MAX_EXCHANGE_BYTES,
        max_unpacked_bytes=128 * 1024 * 1024,
        max_files=256,
    )
    if any(
        b"<!DOCTYPE" in content.replace(b"\0", b"")
        or b"<!ENTITY" in content.replace(b"\0", b"")
        for content in members.values()
    ):
        raise ValueError("XML declarations are not supported")
    for name, content in members.items():
        if name.lower().endswith(".bin") or "/externallinks/" in name.lower():
            raise ValueError("active workbook content is not supported")
        if name.endswith(".rels"):
            try:
                relationships = ElementTree.fromstring(content)
            except ElementTree.ParseError as exc:
                raise ValueError("invalid workbook relationships") from exc
            if any(item.get("TargetMode") == "External" for item in relationships):
                raise ValueError("external workbook relationships are not supported")
    workbook = load_workbook(
        BytesIO(data), read_only=True, data_only=False, keep_links=False
    )
    try:
        if "LX Canonical" not in workbook.sheetnames:
            raise ValueError(
                "workbook has no canonical study exchange; use a column mapping"
            )
        sheet = workbook["LX Canonical"]
        if (
            sheet.max_row is None
            or sheet.max_row > MAX_EXCHANGE_BYTES // _CHUNK_SIZE + 2
            or sheet.max_column != 2
        ):
            raise ValueError("invalid canonical study exchange dimensions")
        rows = sheet.iter_rows()
        first = next(rows, ())
        if [cell.value for cell in first] != [
            "study_exchange_v1",
            "canonical_json_chunks",
        ]:
            raise ValueError("invalid canonical study exchange header")
        chunks: list[str] = []
        size = 0
        for index, row in enumerate(rows):
            if (
                len(row) != 2
                or row[0].value != index
                or row[1].data_type != "s"
                or not isinstance(row[1].value, str)
            ):
                raise ValueError("invalid canonical study exchange chunk")
            size += len(row[1].value.encode())
            if size > MAX_EXCHANGE_BYTES:
                raise ValueError("study exchange exceeds the size limit")
            chunks.append(row[1].value)
        bundle = parse_study_exchange("".join(chunks))
        expected = load_workbook(
            BytesIO(study_exchange_to_xlsx(bundle)),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
        try:
            if workbook.sheetnames != expected.sheetnames:
                raise ValueError(
                    "study exchange worksheets differ from canonical evidence"
                )
            for name in workbook.sheetnames:
                actual_sheet, expected_sheet = workbook[name], expected[name]
                if (actual_sheet.max_row, actual_sheet.max_column) != (
                    expected_sheet.max_row,
                    expected_sheet.max_column,
                ):
                    raise ValueError(
                        "study exchange projection differs from canonical evidence"
                    )
                for actual_row, expected_row in zip_longest(
                    actual_sheet.iter_rows(), expected_sheet.iter_rows(), fillvalue=()
                ):
                    if [(cell.value, cell.data_type) for cell in actual_row] != [
                        (cell.value, cell.data_type) for cell in expected_row
                    ]:
                        raise ValueError(
                            "study exchange projection was edited; use a column mapping to import edits"
                        )
        finally:
            expected.close()
        return bundle
    finally:
        workbook.close()
