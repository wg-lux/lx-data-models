from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook
from pandas.testing import assert_frame_equal

from lx_dtypes.stats.dataset import (
    InterfaceExportDataset,
    KnowledgeBaseDataset,
    LedgerDataset,
)
from lx_dtypes.stats.excel import write_xlsx


def test_excel_export_preserves_literals_identifiers_and_utc(tmp_path: Path) -> None:
    frame = pd.DataFrame(
        {
            "text": ["=1+1", "#N/A", "日本語 – العربية"],
            "identifier": [12345678901234567, 12345678901234568, 12345678901234569],
            "time": pd.date_range("2026-01-02", periods=3, tz="Asia/Tokyo"),
        }
    )
    original = frame.copy(deep=True)
    destination = tmp_path / "export.xlsx"
    write_xlsx(destination, [("data", frame)], overwrite=False)
    workbook = load_workbook(destination)
    try:
        sheet = workbook["data"]
        assert sheet["A2"].value == "=1+1"
        assert sheet["A2"].data_type == "s"
        assert sheet["A3"].value == "#N/A"
        assert sheet["A3"].data_type == "s"
        assert sheet["A4"].value == "日本語 – العربية"
        assert sheet["B2"].value == "12345678901234567"
        assert sheet["C2"].value == pd.Timestamp("2026-01-01 15:00:00")
    finally:
        workbook.close()
    assert_frame_equal(frame, original)


def test_failed_excel_overwrite_preserves_original_and_removes_temporary_file(
    tmp_path: Path,
) -> None:
    destination = tmp_path / "export.xlsx"
    write_xlsx(
        destination, [("original", pd.DataFrame({"value": [1]}))], overwrite=False
    )
    original = destination.read_bytes()
    with pytest.raises(ValueError, match="32767"):
        write_xlsx(
            destination,
            [("data", pd.DataFrame({"value": ["x" * 32768]}))],
            overwrite=True,
        )
    assert destination.read_bytes() == original
    assert list(tmp_path.iterdir()) == [destination]


def test_excel_refuses_overwrite_and_colliding_sheet_names(tmp_path: Path) -> None:
    destination = tmp_path / "export.xlsx"
    destination.write_bytes(b"original")
    with pytest.raises(FileExistsError):
        write_xlsx(destination, [], overwrite=False)
    with pytest.raises(ValueError, match="collide"):
        write_xlsx(
            destination,
            [
                ("data", pd.DataFrame({"value": [1]})),
                ("DATA", pd.DataFrame({"value": [2]})),
            ],
            overwrite=True,
        )
    assert destination.read_bytes() == b"original"


def test_public_dataset_exporters_use_safe_literal_export(tmp_path: Path) -> None:
    # Only the exporter is under test; domain-schema validation has its own lane.
    ledger = LedgerDataset.model_construct(patients=pd.DataFrame({"text": ["=1+1"]}))
    kb = KnowledgeBaseDataset.model_construct(
        citations=pd.DataFrame({"text": ["=1+1"]})
    )
    combined = InterfaceExportDataset.model_construct(ledger=ledger, knowledge_base=kb)
    for name, dataset in (("ledger", ledger), ("kb", kb), ("combined", combined)):
        destination = tmp_path / f"{name}.xlsx"
        dataset.to_xlsx(destination)
        workbook = load_workbook(destination)
        try:
            assert all(sheet["A2"].data_type == "s" for sheet in workbook)
        finally:
            workbook.close()


def test_excel_publish_race_does_not_overwrite_competing_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import os

    destination = tmp_path / "export.xlsx"
    original_link = os.link

    def competing_link(source: Path, target: Path) -> None:
        target.write_bytes(b"concurrent writer")
        original_link(source, target)

    monkeypatch.setattr(os, "link", competing_link)
    with pytest.raises(FileExistsError):
        write_xlsx(
            destination, [("data", pd.DataFrame({"value": [1]}))], overwrite=False
        )
    assert destination.read_bytes() == b"concurrent writer"
    assert list(tmp_path.iterdir()) == [destination]


def test_long_sheet_names_preserve_both_classification_collections(
    tmp_path: Path,
) -> None:
    destination = tmp_path / "export.xlsx"
    write_xlsx(
        destination,
        [
            ("l_p_finding_classification_choices", pd.DataFrame({"value": [1]})),
            (
                "l_p_finding_classification_choice_descriptors",
                pd.DataFrame({"value": [2]}),
            ),
        ],
        overwrite=False,
    )
    workbook = load_workbook(destination)
    try:
        assert len(workbook.sheetnames) == 2
        assert all(len(name) <= 31 for name in workbook.sheetnames)
        assert [sheet["A2"].value for sheet in workbook] == [1, 2]
    finally:
        workbook.close()
