from datetime import UTC, datetime
from io import BytesIO
from uuid import UUID

import pytest
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from pydantic import ValidationError

from lx_dtypes.models.contracts.dtypes_record_persistence import (
    DtypesRecordPersistencePayload,
)
from lx_dtypes.models.contracts.package_snapshot import build_package_snapshot
from lx_dtypes.models.contracts.study import StudyDefinition
from lx_dtypes.models.contracts.study_collection import StudyCollectionRecord
from lx_dtypes.models.contracts.study_exchange import (
    StudyExchangeBundle,
    StudyExchangePatientLink,
    StudyExchangeStudy,
)
from lx_dtypes.models.interface.KnowledgeBase import KnowledgeBase
from lx_dtypes.study_exchange import (
    parse_study_exchange,
    study_exchange_from_xlsx,
    study_exchange_to_xlsx,
)
from lx_dtypes.validation import assess_examination


def bundle() -> StudyExchangeBundle:
    kb = KnowledgeBase.model_validate(
        {
            "config": {"name": "synthetic", "version": "1.0.0"},
            "examination": {
                "synthetic": {"name": "synthetic", "kb_module_name": "synthetic"}
            },
        }
    )
    snapshot = build_package_snapshot(kb)
    record = DtypesRecordPersistencePayload(
        patient=str(UUID(int=4)),
        examination="synthetic",
        knowledge_base_module="synthetic",
        knowledge_base_version="1.0.0",
    )
    study = StudyExchangeStudy(
        study_uuid=UUID(int=1), definition=StudyDefinition(name="=literal")
    )
    collection = StudyCollectionRecord(
        record_uuid=UUID(int=2),
        study_uuid=study.study_uuid,
        study_definition=study.definition,
        patient_examination_uuid=UUID(int=3),
        package_snapshot_id=snapshot.snapshot_id,
        record=record,
        evaluation=assess_examination(kb, record.model_dump()),
        evaluated_at=datetime(2026, 1, 1, tzinfo=UTC),
    )
    return StudyExchangeBundle(
        studies=[study],
        packages=[snapshot],
        collections=[collection],
        patients=[
            StudyExchangePatientLink(
                patient_examination_uuid=UUID(int=3), patient_uuid=UUID(int=4)
            )
        ],
    )


def test_lossless_workbook_preserves_package_and_ledger() -> None:
    original = bundle()
    exported = study_exchange_to_xlsx(original)
    restored = study_exchange_from_xlsx(exported)
    assert restored.model_dump(mode="json") == original.model_dump(mode="json")


@pytest.mark.parametrize("field", ["studies", "packages", "patients"])
def test_missing_closure_rejected(field: str) -> None:
    data = bundle().model_dump(mode="json")
    data[field] = []
    import json

    with pytest.raises(ValidationError):
        parse_study_exchange(json.dumps(data))


def test_rejects_local_patient_identifier_and_mutated_models() -> None:
    original = bundle()
    original.collections[0].record.patient = "123"
    with pytest.raises(ValidationError, match="persistent UUID"):
        study_exchange_to_xlsx(original)


def test_rejects_forged_evaluation() -> None:
    original = bundle()
    # Setting a conflicting evaluated flag is invalid even when no issues exist.
    data = original.model_dump(mode="json")
    data["collections"][0]["evaluation"]["terminology_evaluated"] = False
    import json

    with pytest.raises(ValidationError):
        parse_study_exchange(json.dumps(data))


def test_rejects_formula_in_canonical_chunks() -> None:
    workbook = load_workbook(BytesIO(study_exchange_to_xlsx(bundle())))
    chunk_cell = workbook["LX Canonical"].cell(2, 2)
    assert isinstance(chunk_cell, Cell)
    chunk_cell.value = '=HYPERLINK("https://example.invalid")'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    with pytest.raises(ValueError, match="chunk"):
        study_exchange_from_xlsx(output.getvalue())


def test_empty_study_definition_can_roundtrip() -> None:
    original = StudyExchangeBundle(
        studies=[
            StudyExchangeStudy(
                study_uuid=UUID(int=1), definition=StudyDefinition(name="Empty")
            )
        ]
    )
    assert study_exchange_from_xlsx(study_exchange_to_xlsx(original)).model_dump(
        mode="json"
    ) == original.model_dump(mode="json")


def test_display_values_are_literal_text_and_edits_are_rejected() -> None:
    original = bundle()
    workbook = load_workbook(BytesIO(study_exchange_to_xlsx(original)))
    cell = workbook["Definitions"].cell(2, 4)
    assert isinstance(cell, Cell)
    assert cell.value == "=literal"
    assert cell.data_type == "s"
    cell.value = "Different study name"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    with pytest.raises(ValueError, match="edited"):
        study_exchange_from_xlsx(output.getvalue())


def test_rejects_duplicate_collection_and_unknown_terminology() -> None:
    original = bundle()
    original.collections.append(original.collections[0])
    with pytest.raises(ValidationError, match="duplicate"):
        parse_study_exchange(original)
    original.collections.pop()
    original.collections[0].record.examination = "unknown"
    with pytest.raises(ValidationError, match="assessment"):
        parse_study_exchange(original)


def test_large_canonical_values_are_chunked_without_truncation() -> None:
    original = bundle()
    original.collections[0].record.tags = "x" * 40000
    # Large canonical data is chunked rather than truncated.
    assert study_exchange_from_xlsx(study_exchange_to_xlsx(original)).model_dump(
        mode="json"
    ) == original.model_dump(mode="json")


@pytest.mark.parametrize(
    "value,descriptor_type,cell_type", [(12.5, "numeric", "n"), (True, "boolean", "b")]
)
def test_readable_observations_keep_lesion_identity_and_scalar_types(
    value: float | bool, descriptor_type: str, cell_type: str
) -> None:
    kb = KnowledgeBase.model_validate(
        {
            "config": {"name": "synthetic", "version": "1.0.0"},
            "examination": {
                "synthetic": {
                    "name": "synthetic",
                    "kb_module_name": "synthetic",
                    "findings": ["lesion"],
                }
            },
            "finding": {
                "lesion": {
                    "name": "lesion",
                    "kb_module_name": "synthetic",
                    "classifications": ["size"],
                }
            },
            "classification": {
                "size": {
                    "name": "size",
                    "kb_module_name": "synthetic",
                    "classification_choices": ["measured"],
                }
            },
            "classification_choice": {
                "measured": {
                    "name": "measured",
                    "kb_module_name": "synthetic",
                    "classification_choice_descriptors": ["measurement"],
                }
            },
            "unit": {"mm": {"name": "mm", "kb_module_name": "synthetic"}},
            "classification_choice_descriptor": {
                "measurement": {
                    "name": "measurement",
                    "kb_module_name": "synthetic",
                    "classification_choice_descriptor_type": descriptor_type,
                    "unit": "mm",
                }
            },
        }
    )
    original = bundle()
    payload = original.collections[0].model_dump(mode="python")
    payload["record"]["patient_findings"] = [
        {
            "uuid": str(UUID(int=10 + index)),
            "finding": "lesion",
            "patient_examination": "synthetic",
            "patient_finding_classifications": [
                {
                    "patient_finding": "lesion",
                    "patient_finding_classification_choices": [
                        {
                            "classification": "size",
                            "classification_choice": "measured",
                            "patient_finding_classifications": "size",
                            "patient_finding_classification_choice_descriptors": [
                                {
                                    "classification_choice_descriptor": "measurement",
                                    "patient_finding_classification_choice": "measured",
                                    "descriptor_value": value,
                                }
                            ],
                        }
                    ],
                }
            ],
        }
        for index in range(2)
    ]
    snapshot = build_package_snapshot(kb)
    payload["package_snapshot_id"] = snapshot.snapshot_id
    payload["evaluation"] = assess_examination(kb, payload["record"])
    updated = StudyExchangeBundle(
        studies=original.studies,
        patients=original.patients,
        packages=[snapshot],
        collections=[StudyCollectionRecord.model_validate(payload)],
    )
    workbook = load_workbook(BytesIO(study_exchange_to_xlsx(updated)))
    sheet = workbook["Observations"]
    assert sheet.cell(2, 2).value != sheet.cell(3, 2).value
    assert sheet.cell(2, 12).value == value
    assert sheet.cell(2, 12).data_type == cell_type
    assert sheet.cell(2, 11).value == "mm"
    workbook.close()
